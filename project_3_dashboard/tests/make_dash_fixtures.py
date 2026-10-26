"""Deterministic Project 3 test spreadsheets. Optional manual run (writes into tests\\fixtures):
    python project_3_dashboard\\tests\\make_dash_fixtures.py
Tests build their own copies in temp folders and compute expectations from these row lists.
"""
import csv
import sys
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent / "fixtures"
START = datetime(2025, 1, 1)
CATS = ["Electronics", "Grocery", "Toys"]
REGS = ["North", "South"]
HEADER = ["Date", "Category", "Region", "Units", "Revenue"]
MESSY_HEADER = ["id", "name", "status", "amount", "country", "notes"]
MESSY_ROWS = [(1, "Asha", "paid", 100.5, "IN", None), (2, "Ben", "paid", None, "IN", None),
              (3, "Chen", "pending", 75, "IN", None), (3, "Chen", "pending", 75, "IN", None),
              (4, "Dina", "paid", 20, "IN", None), (5, "Evan", None, 60, "IN", None)]


def orders_rows():
    rows = []
    for i in range(40):
        units = (i * 7) % 13 + 1
        price = 10 + (i % 3) * 5
        rows.append((START + timedelta(days=3 * i), CATS[i % 3], None if i % 13 == 12 else REGS[i % 2],
                     units, None if i % 10 == 9 else round(units * price + 0.5, 2)))
    return rows


def store_b_rows():
    rows = []
    for i in range(12):
        units = (i * 5) % 9 + 2
        price = 12 + (i % 4) * 3
        rows.append((START + timedelta(days=7 * i), CATS[i % 3], REGS[(i + 1) % 2], units,
                     None if i == 5 else round(units * price + 0.25, 2)))
    return rows


def build_orders_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["ACME orders 2025 (internal)"])
    ws.append([])
    ws.append(HEADER)
    for r in orders_rows():
        ws.append(list(r))
    notes = wb.create_sheet("Notes")
    notes.append(["Key", "Value"])
    notes.append(["owner", "finance"])
    notes.append(["currency", "USD"])
    wb.create_sheet("Blank")
    wb.save(path)


def _write_csv(path, header, rows, date_col=None):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in rows:
            w.writerow(["" if v is None else (v.strftime("%Y-%m-%d") if isinstance(v, datetime) else v) for v in r])


def build_all(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    build_orders_xlsx(folder / "orders_2025.xlsx")
    _write_csv(folder / "store_b.csv", HEADER, store_b_rows())
    _write_csv(folder / "messy.csv", MESSY_HEADER, MESSY_ROWS)
    return {"orders": folder / "orders_2025.xlsx", "store_b": folder / "store_b.csv", "messy": folder / "messy.csv"}


if __name__ == "__main__":
    print("Written:", build_all(OUT))
    sys.exit(0)